from __future__ import annotations

import html
from io import BytesIO
from typing import Protocol

import nh3
from docx import Document

from backend.app.domain.document import BlockKind, CanonicalDocument, DocumentBlock


class DigitalExporter(Protocol):
    media_type: str
    extension: str

    def export(self, document: CanonicalDocument) -> bytes:
        ...


def _block_text(document: CanonicalDocument, block: DocumentBlock) -> str:
    return document.accepted_text_for_block(block.id)


class MarkdownExporter:
    media_type = "text/markdown"
    extension = "md"

    def export(self, document: CanonicalDocument) -> bytes:
        lines = [
            f"# {document.source_name}",
            "",
            f"<!-- canonical schema {document.schema_version} -->",
            "",
        ]
        for page in document.pages:
            lines.extend([f"## Page {page.number}", ""])
            for block in page.blocks:
                text = _block_text(document, block)
                if block.kind is BlockKind.heading:
                    lines.extend([f"### {text}", ""])
                elif block.kind is BlockKind.list_item:
                    lines.extend([f"- {text}", ""])
                else:
                    lines.extend([text, ""])
        return "\n".join(lines).encode("utf-8")


class HtmlExporter:
    media_type = "text/html"
    extension = "html"

    def export(self, document: CanonicalDocument) -> bytes:
        parts = [
            "<!doctype html>",
            '<html lang="en"><head><meta charset="utf-8">',
            f"<title>{html.escape(document.source_name)}</title></head><body>",
            '<main data-schema-version="1.0">',
        ]
        for page in document.pages:
            parts.append(f'<section data-page="{page.number}"><h2>Page {page.number}</h2>')
            list_open = False
            for block in page.blocks:
                text = html.escape(_block_text(document, block))
                if block.kind is BlockKind.list_item:
                    if not list_open:
                        parts.append("<ul>")
                        list_open = True
                    parts.append(f"<li>{text}</li>")
                else:
                    if list_open:
                        parts.append("</ul>")
                        list_open = False
                    tag = "h3" if block.kind is BlockKind.heading else "p"
                    parts.append(f"<{tag}>{text}</{tag}>")
            if list_open:
                parts.append("</ul>")
            parts.append("</section>")
        parts.extend(["</main></body></html>"])
        raw = "".join(parts)
        safe = nh3.clean(
            raw,
            tags={
                "html",
                "head",
                "body",
                "meta",
                "title",
                "main",
                "section",
                "h2",
                "h3",
                "p",
                "ul",
                "li",
            },
            attributes={
                "html": {"lang"},
                "main": {"data-schema-version"},
                "section": {"data-page"},
            },
        )
        return safe.encode("utf-8")


class JsonExporter:
    media_type = "application/json"
    extension = "json"

    def export(self, document: CanonicalDocument) -> bytes:
        return (document.model_dump_json(indent=2) + "\n").encode("utf-8")


class DocxExporter:
    media_type = (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    extension = "docx"

    def export(self, document: CanonicalDocument) -> bytes:
        output = Document()
        output.core_properties.title = document.source_name
        output.core_properties.subject = "OCR canonical document"
        for page in document.pages:
            output.add_heading(f"Page {page.number}", level=1)
            for block in page.blocks:
                text = _block_text(document, block)
                if block.kind is BlockKind.heading:
                    output.add_heading(text, level=2)
                elif block.kind is BlockKind.list_item:
                    output.add_paragraph(text, style="List Bullet")
                else:
                    output.add_paragraph(text)
        stream = BytesIO()
        output.save(stream)
        return stream.getvalue()


DIGITAL_EXPORTERS: dict[str, DigitalExporter] = {
    "markdown": MarkdownExporter(),
    "html": HtmlExporter(),
    "json": JsonExporter(),
    "docx": DocxExporter(),
}
